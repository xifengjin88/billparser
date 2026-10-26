"""Generate test files: known provider exports plus messy custom layouts.

    python samples/make_samples.py        -> samples/files/*

Column names for provider exports follow public documentation; all values are synthetic.
"""

from __future__ import annotations

import gzip
import io
import json
import random
import re
import shutil
import subprocess
import tempfile
from datetime import date, timedelta
from pathlib import Path

import pandas as pd

OUT = Path(__file__).parent / "files"
rnd = random.Random(42)
DAYS = [date(2026, 9, 1) + timedelta(days=i) for i in range(30)]


def aws_cur2() -> None:
    items = [("AmazonEC2", "Amazon Elastic Compute Cloud", "USE1-BoxUsage:m5.xlarge", "Hrs", 0.192),
             ("AmazonEC2", "Amazon Elastic Compute Cloud", "USE1-EBS:VolumeUsage.gp3", "GB-Mo", 0.08 / 30),
             ("AmazonRDS", "Amazon Relational Database Service", "USE1-InstanceUsage:db.r6g.large", "Hrs", 0.25),
             ("AmazonS3", "Amazon Simple Storage Service", "USE1-TimedStorage-ByteHrs", "GB-Mo", 0.023 / 30),
             ("AWSLambda", "AWS Lambda", "USE1-Request", "Requests", 0.0000002)]
    rows = []
    for d in DAYS:
        for code, name, ut, unit, rate in items:
            for _ in range(rnd.randint(8, 14)):
                q = {"Hrs": 24, "GB-Mo": rnd.uniform(50, 500), "Requests": rnd.randint(10**5, 10**6)}[unit]
                rows.append(_cur_row(d, code, name, ut, unit, rate, q, "Usage"))
    rows.append(_cur_row(DAYS[-1], "AmazonEC2", "Amazon Elastic Compute Cloud", "Tax", "", 0, 0, "Tax", cost=412.55))
    rows.append(_cur_row(DAYS[-1], "AmazonEC2", "Amazon Elastic Compute Cloud", "Credit", "", 0, 0, "Credit", cost=-150.0))
    buf = io.StringIO()
    pd.DataFrame(rows).to_csv(buf, index=False)
    (OUT / "aws_cur2.csv.gz").write_bytes(gzip.compress(buf.getvalue().encode()))


def _cur_row(d, code, name, ut, unit, rate, q, typ, cost=None):
    return {"bill_invoice_id": "EUINUS26-123456", "bill_billing_period_start_date": "2026-09-01T00:00:00Z",
            "line_item_usage_account_id": "111122223333", "line_item_line_item_type": typ,
            "line_item_usage_start_date": f"{d}T00:00:00Z", "line_item_usage_end_date": f"{d + timedelta(days=1)}T00:00:00Z",
            "line_item_product_code": code, "line_item_usage_type": ut, "line_item_usage_amount": round(q, 6),
            "line_item_unblended_rate": rate, "line_item_unblended_cost": round(cost if cost is not None else q * rate, 8),
            "line_item_currency_code": "USD", "line_item_line_item_description": f"{ut} usage",
            "product_product_name": name, "product_region_code": "us-east-1", "pricing_unit": unit,
            "pricing_public_on_demand_cost": round(cost if cost is not None else q * rate, 8)}


CUR_RI = "arn:aws:ec2:us-east-1:111122223333:reserved-instances/1d3fbc13-f181-4c40-9dd6-12b345678de9"
CUR_SP = "arn:aws:savingsplans::111122223333:savingsplan/7a2c1b9e-0d55-4f0e-9c1a-3b5e2f8d1c44"


def _cur_commitment_rows() -> list[dict]:
    """Three days of one account in us-east-1: on-demand, Spot, Reserved Instance and Savings Plan usage, the
    negation of the covered usage, the commitment fees, a tax and a credit."""
    ec2, rds = ("AmazonEC2", "Amazon Elastic Compute Cloud"), ("AmazonRDS", "Amazon Relational Database Service")
    rows = []
    for d in DAYS[:3]:
        rows.append(_cur_row(d, *ec2, "USE1-BoxUsage:m5.xlarge", "Hrs", 0.192, 24, "Usage"))
        rows.append(_cur_row(d, *ec2, "USE1-SpotUsage:m5.xlarge", "Hrs", 0.06, 24, "Usage"))
        ri = _cur_row(d, *rds, "USE1-InstanceUsage:db.r6g.large", "Hrs", 0.0, 24, "DiscountedUsage")
        ri.update(pricing_public_on_demand_cost=round(24 * 0.25, 8), reservation_reservation_a_r_n=CUR_RI,
                  reservation_effective_cost=round(24 * 0.15, 8))
        rows.append(ri)
        sp = _cur_row(d, *ec2, "USE1-BoxUsage:c5.large", "Hrs", 0.085, 24, "SavingsPlanCoveredUsage")
        sp.update(savings_plan_savings_plan_a_r_n=CUR_SP, savings_plan_savings_plan_effective_cost=round(24 * 0.06, 8))
        rows.append(sp)
        neg = _cur_row(d, *ec2, "USE1-BoxUsage:c5.large", "Hrs", 0.085, 24, "SavingsPlanNegation",
                       cost=-round(24 * 0.085, 8))
        neg.update(savings_plan_savings_plan_a_r_n=CUR_SP, pricing_public_on_demand_cost=None)
        rows.append(neg)
    last = DAYS[2]
    fee = _cur_row(DAYS[0], "ComputeSavingsPlans", "Savings Plans for AWS Compute usage", "ComputeSP:1yrNoUpfront",
                   "", 0, 0, "SavingsPlanRecurringFee", cost=round(3 * 24 * 0.06, 8))
    fee.update(savings_plan_savings_plan_a_r_n=CUR_SP, pricing_public_on_demand_cost=None)
    rifee = _cur_row(DAYS[0], *rds, "USE1-HeavyUsage:db.r6g.large", "Hrs", 0, 0, "RIFee", cost=round(3 * 24 * 0.15, 8))
    rifee.update(reservation_reservation_a_r_n=CUR_RI, pricing_public_on_demand_cost=None)
    rows += [fee, rifee,
             _cur_row(last, *ec2, "Tax", "", 0, 0, "Tax", cost=2.50),
             _cur_row(last, *ec2, "Credit", "", 0, 0, "Credit", cost=-1.00)]
    return rows


def _write_cur(rows: list[dict], name: str) -> float:
    buf = io.StringIO()
    pd.DataFrame(rows).to_csv(buf, index=False)
    (OUT / name).write_bytes(gzip.compress(buf.getvalue().encode()))
    return round(sum(r["line_item_unblended_cost"] for r in rows), 2)


def aws_cur2_commitments() -> float:
    """CUR 2.0 with commitments and the amortized (effective cost) columns. Returns the billed total."""
    return _write_cur(_cur_commitment_rows(), "aws_cur2_commitments.csv.gz")


def aws_cur2_no_amortized() -> float:
    """The same rows without the reservation and Savings Plan effective-cost columns."""
    rows = [{k: v for k, v in r.items() if k not in ("reservation_effective_cost",
                                                     "savings_plan_savings_plan_effective_cost")}
            for r in _cur_commitment_rows()]
    return _write_cur(rows, "aws_cur2_no_amortized.csv.gz")


# legacy CUR names (a real export's header): the instance type, engine, SKU and rate code by name; the region code
# empty on data transfer; unused commitment on the RI and Savings Plan fee lines; product columns the spec does not
# name (product/sku must never be read as an instance type)
LEGACY_SKU = {"ec2": "RVS26QBX9XP34QN6", "dt": "AQ7UHUN3VCX2B3GM", "rds": "6WSK9HTNPZXDNS5A", "c5": "JRTCKXETXFVB2RYV"}


def _legacy_rows() -> list[dict]:
    def row(d, typ, code, name, ut, unit, q, rate, cost, sku, **extra) -> dict:
        r = {"identity/LineItemId": f"li-{len(rows)}", "bill/InvoiceId": "EUINUS26-654321",
             "bill/PayerAccountId": "111122223333", "lineItem/UsageAccountId": "444455556666",
             "lineItem/LineItemType": typ, "lineItem/UsageStartDate": f"{d}T00:00:00Z",
             "lineItem/UsageEndDate": f"{d + timedelta(days=1)}T00:00:00Z", "lineItem/ProductCode": code,
             "lineItem/UsageType": ut, "lineItem/UsageAmount": q, "lineItem/CurrencyCode": "USD",
             "lineItem/UnblendedRate": rate, "lineItem/UnblendedCost": round(cost, 8),
             "lineItem/LineItemDescription": f"{ut} {typ}", "product/ProductName": name,
             "product/region": "us-east-1", "product/regionCode": "us-east-1", "product/sku": sku,
             "product/instancesku": "", "product/vcpu": "", "pricing/RateCode": f"{sku}.JRTCKXETXF.6YS6EN2CT7" if sku else "",
             "pricing/term": "OnDemand", "pricing/unit": unit, "pricing/publicOnDemandCost": round(cost, 8),
             "resourceTags/user:Name": "web"}
        r.update(extra)
        return r
    rows: list[dict] = []
    ec2, rds = ("AmazonEC2", "Amazon Elastic Compute Cloud"), ("AmazonRDS", "Amazon Relational Database Service")
    for d in DAYS[:2]:
        rows.append(row(d, "Usage", *ec2, "USE1-BoxUsage:m5.xlarge", "Hrs", 24, 0.192, 24 * 0.192, LEGACY_SKU["ec2"],
                        **{"product/instanceType": "m5.xlarge", "product/operatingSystem": "Linux", "product/vcpu": "4"}))
        rows.append(row(d, "Usage", *ec2, "USE1-USW2-AWS-Out-Bytes", "GB", 10, 0.02, 0.2, LEGACY_SKU["dt"],
                        **{"product/regionCode": "", "product/region": "us-west-2"}))
        rows.append(row(d, "DiscountedUsage", *rds, "USE1-InstanceUsage:db.r6g.large", "Hrs", 24, 0.0, 0.0,
                        LEGACY_SKU["rds"], **{"product/instanceType": "db.r6g.large",
                                              "product/databaseEngine": "PostgreSQL",
                                              "product/deploymentOption": "Multi-AZ", "pricing/term": "Reserved",
                                              "pricing/publicOnDemandCost": 6.0, "reservation/ReservationARN": CUR_RI,
                                              "reservation/EffectiveCost": 3.6}))
        rows.append(row(d, "SavingsPlanCoveredUsage", *ec2, "USE1-BoxUsage:c5.large", "Hrs", 24, 0.085, 2.04,
                        LEGACY_SKU["c5"], **{"product/instanceType": "c5.large", "savingsPlan/SavingsPlanARN": CUR_SP,
                                             "savingsPlan/SavingsPlanEffectiveCost": 1.44}))
        rows.append(row(d, "SavingsPlanNegation", *ec2, "USE1-BoxUsage:c5.large", "Hrs", 24, 0.085, -2.04,
                        LEGACY_SKU["c5"], **{"savingsPlan/SavingsPlanARN": CUR_SP, "pricing/RateCode": "",
                                             "pricing/publicOnDemandCost": ""}))
    # the RI fee: 10.80 of commitment, 7.20 used by the two days above, 3.60 unused (1.00 upfront + 2.60 recurring)
    rows.append(row(DAYS[0], "RIFee", *rds, "USE1-HeavyUsage:db.r6g.large", "Hrs", 72, 0.15, 10.8, LEGACY_SKU["ec2"],
                    **{"reservation/ReservationARN": CUR_RI, "reservation/UnusedAmortizedUpfrontFeeForBillingPeriod": 1.0,
                       "reservation/UnusedRecurringFee": 2.6, "pricing/term": "Reserved",
                       "pricing/publicOnDemandCost": 0}))
    # the Savings Plan fee: 4.32 committed to date, 2.88 used
    rows.append(row(DAYS[0], "SavingsPlanRecurringFee", "ComputeSavingsPlans", "Savings Plans for AWS Compute usage",
                    "ComputeSP:1yrNoUpfront", "", 0, 0, 4.32, "N3RQ9N6KR9HDPCZ9",
                    **{"savingsPlan/SavingsPlanARN": CUR_SP, "savingsPlan/TotalCommitmentToDate": 4.32,
                       "savingsPlan/UsedCommitment": 2.88, "pricing/publicOnDemandCost": 0}))
    rows.append(row(DAYS[1], "Tax", *ec2, "", "", 0, 0, 1.5, "", **{"product/regionCode": "", "product/region": ""}))
    return rows


def aws_cur_legacy() -> float:
    """A legacy CUR (lineItem/... names) with commitments and unused commitment on the fee lines."""
    rows = _legacy_rows()
    cols = list(dict.fromkeys(k for r in rows for k in r))
    buf = io.StringIO()
    pd.DataFrame(rows, columns=cols).to_csv(buf, index=False)
    (OUT / "aws_cur_legacy.csv.gz").write_bytes(gzip.compress(buf.getvalue().encode()))
    return round(sum(r["lineItem/UnblendedCost"] for r in rows), 2)


def aws_cur2_product_map() -> float:
    """CUR 2.0 whose engine, OS, deployment and product name are in the "product" map column (JSON text), with the
    SKU, rate code, unused commitment columns and a usage account name."""
    def snake(h: str) -> str:
        group, name = h.split("/", 1)
        group = {"lineItem": "line_item", "savingsPlan": "savings_plan"}.get(group, group)
        return f"{group}_" + re.sub(r"(?<=[a-z0-9])([A-Z])", r"_\1", name).lower()
    in_map = {"product/ProductName": "product_name", "product/operatingSystem": "operating_system",
              "product/databaseEngine": "database_engine", "product/deploymentOption": "deployment_option",
              "product/region": "region", "product/vcpu": "vcpu"}
    out = []
    for r in _legacy_rows():
        o = {}
        for k, v in r.items():
            if k in in_map or k.startswith("resourceTags/") or k == "product/instancesku":
                continue
            o[{"product/regionCode": "product_region_code", "product/instanceType": "product_instance_type",
               "reservation/ReservationARN": "reservation_reservation_a_r_n",
               "savingsPlan/SavingsPlanARN": "savings_plan_savings_plan_a_r_n",
               "pricing/RateCode": "pricing_rate_code", "identity/LineItemId": "identity_line_item_id",
               "pricing/publicOnDemandCost": "pricing_public_on_demand_cost",
               "savingsPlan/SavingsPlanEffectiveCost": "savings_plan_savings_plan_effective_cost",
               }.get(k, snake(k))] = v
        o["product"] = json.dumps({m: r[k] for k, m in in_map.items() if r.get(k) not in (None, "")})
        o["line_item_usage_account_name"] = "prod"
        out.append(o)
    cols = list(dict.fromkeys(k for r in out for k in r))
    buf = io.StringIO()
    pd.DataFrame(out, columns=cols).to_csv(buf, index=False)
    (OUT / "aws_cur2_product_map.csv.gz").write_bytes(gzip.compress(buf.getvalue().encode()))
    return round(sum(r["line_item_unblended_cost"] for r in out), 2)


# ---------------------------------------------------------------- task 03: every attribute and ID column filled
FULL_LEGACY = {"lineItem/Operation": "RunInstances", "product/tenancy": "Shared",
               "product/licenseModel": "No License required", "pricing/PurchaseOption": "",
               "lineItem/ResourceId": "i-0abc123def4567890"}


def _legacy_full_rows() -> list[dict]:
    """The legacy CUR rows of aws_cur_legacy, with operation, tenancy, license model, public on-demand rate, purchase
    option and resource ID filled where a real CUR fills them."""
    rows = []
    for r in _legacy_rows():
        r = dict(r, **FULL_LEGACY)
        ut, typ = r["lineItem/UsageType"], r["lineItem/LineItemType"]
        q = float(r["lineItem/UsageAmount"] or 0)
        r["pricing/publicOnDemandRate"] = round(float(r["pricing/publicOnDemandCost"] or 0) / q, 6) if q else 0
        if "InstanceUsage" in ut or "HeavyUsage" in ut:
            r.update({"lineItem/Operation": "CreateDBInstance:0014", "product/licenseModel": "No license required",
                      "lineItem/ResourceId": "arn:aws:rds:us-east-1:444455556666:db:orders"})
        if "Out-Bytes" in ut:
            r.update({"lineItem/Operation": "", "product/tenancy": "", "product/licenseModel": "",
                      "lineItem/ResourceId": ""})
        if typ in ("RIFee", "DiscountedUsage"):
            r["pricing/PurchaseOption"] = "No Upfront"
        if typ in ("SavingsPlanRecurringFee", "Tax"):
            r.update({"lineItem/Operation": "", "product/tenancy": "", "product/licenseModel": "",
                      "lineItem/ResourceId": ""})
        if "c5.large" in ut:
            r["product/operatingSystem"] = "Linux"
        if "BoxUsage" in ut or "InstanceUsage" in ut or "HeavyUsage" in ut:
            r.setdefault("product/instanceType", "")
        rows.append(r)
    # the RI fee is for the RDS instance: its type and engine
    for r in rows:
        if r["lineItem/LineItemType"] == "RIFee":
            r.update({"product/instanceType": "db.r6g.large", "product/databaseEngine": "PostgreSQL",
                      "product/deploymentOption": "Multi-AZ"})
    return rows


def _write_gz_csv(rows: list[dict], name: str) -> None:
    cols = list(dict.fromkeys(k for r in rows for k in r))
    buf = io.StringIO()
    pd.DataFrame(rows, columns=cols).to_csv(buf, index=False)
    (OUT / name).write_bytes(gzip.compress(buf.getvalue().encode()))


def aws_cur_legacy_full() -> float:
    rows = _legacy_full_rows()
    _write_gz_csv(rows, "aws_cur_legacy_full.csv.gz")
    return round(sum(r["lineItem/UnblendedCost"] for r in rows), 2)


CUR2_MAP_KEYS = {"product/ProductName": "product_name", "product/operatingSystem": "operating_system",
                 "product/databaseEngine": "database_engine", "product/deploymentOption": "deployment_option",
                 "product/region": "region", "product/tenancy": "tenancy", "product/licenseModel": "license_model"}


def _cur2_name(h: str) -> str:
    """A legacy CUR header as its CUR 2.0 column (lineItem/UsageAccountId -> line_item_usage_account_id)."""
    special = {"product/regionCode": "product_region_code", "product/instanceType": "product_instance_type",
               "reservation/ReservationARN": "reservation_reservation_a_r_n",
               "savingsPlan/SavingsPlanARN": "savings_plan_savings_plan_a_r_n", "pricing/RateCode": "pricing_rate_code",
               "pricing/PurchaseOption": "pricing_purchase_option", "identity/LineItemId": "identity_line_item_id",
               "savingsPlan/SavingsPlanEffectiveCost": "savings_plan_savings_plan_effective_cost"}
    if h in special:
        return special[h]
    group, name = h.split("/", 1)
    group = {"lineItem": "line_item", "savingsPlan": "savings_plan"}.get(group, group)
    return f"{group}_" + re.sub(r"(?<=[a-z0-9])([A-Z])", r"_\1", name).lower()


def _cur2_full_rows(as_map: bool) -> list[dict]:
    out = []
    for r in _legacy_full_rows():
        o = {}
        for k, v in r.items():
            if k.startswith("resourceTags/") or k in ("product/instancesku", "product/vcpu"):
                continue
            if as_map and k in CUR2_MAP_KEYS:
                continue
            o[_cur2_name(k) if not (k in CUR2_MAP_KEYS) else "product_" + CUR2_MAP_KEYS[k]] = v
        if as_map:
            o["product"] = json.dumps({m: r[k] for k, m in CUR2_MAP_KEYS.items() if r.get(k) not in (None, "")})
        out.append(o)
    return out


def aws_cur2_flat() -> float:
    """CUR 2.0 with every product attribute as its own column (an export that selects product.x AS product_x)."""
    rows = _cur2_full_rows(as_map=False)
    _write_gz_csv(rows, "aws_cur2_flat.csv.gz")
    return round(sum(r["line_item_unblended_cost"] for r in rows), 2)


def aws_cur2_map() -> float:
    """CUR 2.0 with the product attributes in the product map column (JSON text)."""
    rows = _cur2_full_rows(as_map=True)
    _write_gz_csv(rows, "aws_cur2_map.csv.gz")
    return round(sum(r["line_item_unblended_cost"] for r in rows), 2)


def azure_usage_full() -> float:
    """Azure cost details (MCA, amortized) with MeterId, ReservationId, BenefitId, PayGPrice, UnitPrice, ResourceId and
    AdditionalInfo; VM meter names are sizes (D4s v5), one of them Spot."""
    rid, sp = "11111111-2222-3333-4444-555555555555", "/providers/Microsoft.BillingBenefits/savingsPlanOrders/abc/savingsPlans/def"
    vm = "/subscriptions/aaaaaaaa-1111-2222-3333-444444444444/resourceGroups/web/providers/Microsoft.Compute/virtualMachines/"
    rows = []

    def row(d, cat, meter, model, q, payg, unit, eff, meter_id, info=None, **extra):
        return {"SubscriptionId": "aaaaaaaa-1111-2222-3333-444444444444", "SubscriptionName": "Prod",
                "Date": d.strftime("%m/%d/%Y"), "MeterCategory": cat, "MeterSubCategory": "", "MeterName": meter,
                "MeterId": meter_id, "ProductId": "DZH318Z0BQ4L", "Quantity": q, "PayGPrice": payg, "UnitPrice": unit,
                "EffectivePrice": eff, "CostInBillingCurrency": round(q * eff, 6), "UnitOfMeasure": "1 Hour",
                "ChargeType": "Usage", "PricingModel": model, "ResourceLocation": "eastus", "BillingCurrencyCode": "USD",
                "AdditionalInfo": json.dumps(info) if info else "", **extra}
    for d in DAYS[:3]:
        rows.append(row(d, "Virtual Machines", "D4s v5", "OnDemand", 24, 0.20, 0.18, 0.18,
                        "0f1e2d3c-0000-1111-2222-333344445555", {"ServiceType": "Standard_D4s_v5", "VCPUs": 4},
                        ResourceId=vm + "web-1"))
        rows.append(row(d, "Virtual Machines", "D4s v5 Spot", "Spot", 24, 0.20, 0.05, 0.05,
                        "0f1e2d3c-0000-1111-2222-333344445556", ResourceId=vm + "batch-1"))
        rows.append(row(d, "Virtual Machines", "E8ds v5", "Reservation", 24, 0.50, 0.45, 0.30,
                        "0f1e2d3c-0000-1111-2222-333344445557", {"ServiceType": "Standard_E8ds_v5"}, ReservationId=rid,
                        ResourceId=vm + "db-1"))
        rows.append(row(d, "Virtual Machines", "B2ms", "SavingsPlan", 24, 0.08, 0.08, 0.06,
                        "0f1e2d3c-0000-1111-2222-333344445558", BenefitId=sp, ResourceId=vm + "tools-1"))
        rows.append(row(d, "Storage", "P10 Disks", "OnDemand", 1, 0.66, 0.66, 0.66,
                        "0f1e2d3c-0000-1111-2222-333344445559",
                        ResourceId="/subscriptions/aaaaaaaa-1111-2222-3333-444444444444/resourceGroups/web/providers/"
                                   "Microsoft.Compute/disks/web-1-os"))
    pd.DataFrame(rows).to_csv(OUT / "azure_usage_full.csv", index=False)
    return round(sum(r["CostInBillingCurrency"] for r in rows), 2)


def focus_full() -> float:
    """FOCUS with the SKU, SKU price, resource, commitment and unit-price columns."""
    rows = []

    def row(d, svc, sku, billed, list_p, contracted_p, q, pricing, cid=None, res=None):
        return {"BilledCost": billed, "EffectiveCost": billed, "ListCost": round(list_p * q, 6),
                "ContractedCost": round(contracted_p * q, 6), "ListUnitPrice": list_p, "ContractedUnitPrice": contracted_p,
                "BillingCurrency": "USD", "ChargePeriodStart": f"{d}T00:00:00Z",
                "ChargePeriodEnd": f"{d + timedelta(days=1)}T00:00:00Z", "ChargeCategory": "Usage",
                "PricingCategory": pricing, "CommitmentDiscountId": cid,
                "CommitmentDiscountType": "Savings Plan" if cid else None, "ProviderName": "AWS", "ServiceName": svc,
                "SkuId": sku, "SkuPriceId": f"{sku}.JRTCKXETXF.6YS6EN2CT7", "ResourceId": res,
                "SubAccountId": "111122223333", "RegionId": "us-east-1", "ConsumedQuantity": q, "ConsumedUnit": "Hours",
                "ChargeDescription": f"${list_p} per On Demand Linux {'m5.xlarge' if pricing == 'Standard' else 'c5.large'} "
                                     "Instance Hour",
                "x_UsageType": f"USE1-BoxUsage:{'m5.xlarge' if pricing == 'Standard' else 'c5.large'}",
                "x_Operation": "RunInstances"}
    for d in DAYS[:3]:
        rows.append(row(d, "Amazon EC2", "RVS26QBX9XP34QN6", 4.608, 0.192, 0.192, 24, "Standard",
                        res="arn:aws:ec2:us-east-1:111122223333:instance/i-0abc"))
        rows.append(row(d, "Amazon EC2", "JRTCKXETXFVB2RYV", 1.44, 0.085, 0.085, 24, "Committed",
                        cid="arn:aws:savingsplans::111122223333:savingsplan/7a2c", res="arn:aws:ec2:us-east-1:111122223333:instance/i-0def"))
    pd.DataFrame(rows).to_parquet(OUT / "focus_full.parquet", index=False)
    return round(sum(r["BilledCost"] for r in rows), 2)


def gcp_export_full() -> float:
    """GCP detailed export with sku.id, resource.name and system labels (the machine type)."""
    recs = []
    for d in DAYS[:3]:
        for sku_id, sku, cost, spec, name in [
                ("CP-COMPUTEENGINE-VMIMAGE-N2-STANDARD-4", "N2 Instance Core running in Americas", 50.0, "n2-standard-4", "web-1"),
                ("D973-5D65-BAB2", "E2 Instance Core running in Americas", 12.0, "e2-medium", "tools-1")]:
            recs.append({"billing_account_id": "01ABCD-234567-89EFGH",
                         "service": {"id": "6F81-5844-456A", "description": "Compute Engine"},
                         "sku": {"id": sku_id, "description": sku},
                         "usage_start_time": f"{d}T00:00:00Z", "usage_end_time": f"{d}T23:59:59Z",
                         "project": {"id": "shop-prod", "name": "shop-prod"}, "location": {"region": "us-east1"},
                         "resource": {"name": f"projects/shop-prod/zones/us-east1-b/instances/{name}",
                                      "global_name": f"//compute.googleapis.com/projects/shop-prod/instances/{name}"},
                         "system_labels": [{"key": "compute.googleapis.com/machine_spec", "value": spec},
                                           {"key": "compute.googleapis.com/cores", "value": "4"}],
                         "cost": cost, "currency": "USD", "usage": {"amount_in_pricing_units": 24, "pricing_unit": "hour"},
                         "credits": [], "cost_type": "regular"})
    (OUT / "gcp_export_full.json").write_text(json.dumps(recs))
    return round(sum(r["cost"] for r in recs), 2)


def _azure_rows(amortized: bool) -> list[dict]:
    """Azure cost details (EA) in one export style: on-demand, Spot and reservation-covered VM usage, and the
    reservation's purchase (actual) or its unused part (both styles; it costs 0 in the actual export)."""
    rid = "11111111-2222-3333-4444-555555555555"
    rows = []

    def row(d, model, charge, meter, q, payg, unit, eff, **extra):
        return {"SubscriptionId": "aaaaaaaa-1111-2222-3333-444444444444", "SubscriptionName": "Prod",
                "Date": d.strftime("%m/%d/%Y"), "MeterCategory": "Virtual Machines", "MeterName": meter,
                "ProductId": "DZH318Z0BQ4L", "Quantity": q, "PayGPrice": payg, "UnitPrice": unit,
                "EffectivePrice": eff, "CostInBillingCurrency": round(q * eff, 6), "UnitOfMeasure": "1 Hour",
                "ChargeType": charge, "PricingModel": model, "ResourceLocation": "eastus",
                "BillingCurrencyCode": "USD", **extra}
    for d in DAYS[:3]:
        rows.append(row(d, "OnDemand", "Usage", "D4s v5", 24, 0.20, 0.18, 0.18))
        rows.append(row(d, "Spot", "Usage", "D4s v5 Spot", 24, 0.20, 0.05, 0.05))
        rows.append(row(d, "Reservation", "Usage", "D8s v5", 24, 0.40, 0.36, 0.22 if amortized else 0.0,
                        ReservationId=rid))
    rows.append(row(DAYS[2], "Reservation", "UnusedReservation", "D8s v5", 2, 0.40, 0.36, 0.22 if amortized else 0.0,
                    ReservationId=rid))
    if not amortized:
        rows.append(row(DAYS[0], "Reservation", "Purchase", "Reserved VM Instance, D8s v5, 1 Year", 1, 160.6, 160.6, 160.6,
                        ReservationId=rid))
    return rows


def azure_amortized() -> float:
    rows = _azure_rows(amortized=True)
    pd.DataFrame(rows).to_csv(OUT / "azure_amortized.csv", index=False)
    return round(sum(r["CostInBillingCurrency"] for r in rows), 2)


def azure_actual() -> float:
    rows = _azure_rows(amortized=False)
    pd.DataFrame(rows).to_csv(OUT / "azure_actual.csv", index=False)
    return round(sum(r["CostInBillingCurrency"] for r in rows), 2)


def focus_commitments() -> float:
    """FOCUS with every cost type and the commitment discount columns: on-demand, Savings Plan, Reserved Instance,
    unused commitment, Spot, a commitment purchase and a tax."""
    rows = []

    def row(d, cat, pricing, billed, listed, contracted, effective, ctype=None, status=None, svc="Amazon EC2"):
        return {"BilledCost": billed, "ListCost": listed, "ContractedCost": contracted, "EffectiveCost": effective,
                "BillingCurrency": "USD", "ChargePeriodStart": f"{d}T00:00:00Z",
                "ChargePeriodEnd": f"{d + timedelta(days=1)}T00:00:00Z", "ChargeCategory": cat,
                "PricingCategory": pricing, "CommitmentDiscountId": f"arn:{ctype}" if ctype else None,
                "CommitmentDiscountType": ctype, "CommitmentDiscountStatus": status,
                "ProviderName": "AWS", "ServiceName": svc, "SubAccountId": "111122223333", "RegionId": "us-east-1",
                "ConsumedQuantity": 24.0, "ConsumedUnit": "Hours"}
    for d in DAYS[:3]:
        rows.append(row(d, "Usage", "Standard", 4.32, 4.80, 4.32, 4.32))
        rows.append(row(d, "Usage", "Committed", 0.0, 2.04, 1.94, 1.44, "Savings Plan", "Used"))
        rows.append(row(d, "Usage", "Committed", 0.0, 6.00, 5.70, 3.60, "Reserved Instance", "Used", "Amazon RDS"))
        rows.append(row(d, "Usage", "Dynamic", 1.44, 4.61, 1.44, 1.44))
    rows.append(row(DAYS[2], "Usage", "Committed", 0.0, 0.0, 0.0, 0.30, "Savings Plan", "Unused"))
    rows.append(row(DAYS[0], "Purchase", None, 4.32, 4.32, 4.32, 0.0, "Savings Plan"))
    rows.append(row(DAYS[2], "Tax", None, 1.25, 1.25, 1.25, 1.25, svc="Tax"))
    pd.DataFrame(rows).to_parquet(OUT / "focus_commitments.parquet", index=False)
    return round(sum(r["BilledCost"] for r in rows), 2)


def gcp_commitments() -> float:
    """GCP billing export with a committed use discount credit on some usage, and cost_at_list."""
    recs = []
    for d in DAYS[:3]:
        for sku, cost, credits in [
                ("N2 Instance Core running in Americas", 50.0,
                 [{"name": "Committed use discount: N2 Core", "amount": -20.0, "type": "COMMITTED_USAGE_DISCOUNT"}]),
                ("Spot Preemptible N2 Instance Core running in Americas", 8.0, []),
                ("E2 Instance Core running in Americas", 12.0,
                 [{"name": "Sustained use discount", "amount": -2.4, "type": "SUSTAINED_USAGE_DISCOUNT"}])]:
            recs.append({"billing_account_id": "01ABCD-234567-89EFGH",
                         "service": {"id": "6F81-5844-456A", "description": "Compute Engine"},
                         "sku": {"id": "AAAA-BBBB", "description": sku},
                         "usage_start_time": f"{d}T00:00:00Z", "usage_end_time": f"{d}T23:59:59Z",
                         "project": {"id": "shop-prod", "name": "shop-prod"}, "location": {"region": "us-east1"},
                         "cost": cost, "cost_at_list": round(cost * 1.1, 2), "currency": "USD",
                         "usage": {"amount_in_pricing_units": 24, "pricing_unit": "hour"},
                         "credits": credits, "cost_type": "regular"})
    recs.append({"billing_account_id": "01ABCD-234567-89EFGH", "service": {"id": "", "description": "Tax"},
                 "sku": {"id": "", "description": "Tax"}, "usage_start_time": f"{DAYS[2]}T00:00:00Z",
                 "usage_end_time": f"{DAYS[2]}T23:59:59Z", "project": {"id": "shop-prod", "name": "shop-prod"},
                 "location": {"region": ""}, "cost": 3.0, "currency": "USD", "credits": [], "cost_type": "tax"})
    (OUT / "gcp_commitments.json").write_text(json.dumps(recs))
    return round(sum(r["cost"] for r in recs), 2)


def azure_ea() -> None:
    meters = [("Virtual Machines", "D4s v5", "DZH318Z0BQ4L", "1 Hour", 0.192),
              ("Storage", "P10 Disks", "DZH318Z0BQ8K", "1/Month", 19.71 / 30),
              ("SQL Database", "vCore", "DZH318Z0BQ1C", "1 Hour", 0.505),
              ("Azure Kubernetes Service", "Standard Uptime SLA", "DZH318Z0BQ3V", "1 Hour", 0.10),
              ("Bandwidth", "Standard Data Transfer Out", "DZH318Z0BRS8", "1 GB", 0.087)]
    rows = []
    for d in DAYS:
        for cat, meter, pid, uom, price in meters:
            for _ in range(rnd.randint(5, 10)):
                q = 24 if "Hour" in uom else 1 if "Month" in uom else rnd.uniform(1, 80)
                rows.append({"BillingAccountId": "12345678", "BillingAccountName": "Contoso", "BillingPeriodStartDate": "09/01/2026",
                             "BillingPeriodEndDate": "09/30/2026", "SubscriptionId": "aaaaaaaa-1111-2222-3333-444444444444",
                             "SubscriptionName": "Prod", "Date": d.strftime("%m/%d/%Y"), "ProductId": pid,
                             "MeterId": "bbbbbbbb-1111-2222-3333-444444444444", "MeterName": meter, "MeterCategory": cat,
                             "MeterRegion": "US East", "Quantity": round(q, 4), "EffectivePrice": price,
                             "CostInBillingCurrency": round(q * price, 6), "UnitOfMeasure": uom, "ChargeType": "Usage",
                             "ResourceLocation": "eastus", "BillingCurrencyCode": "USD"})
    pd.DataFrame(rows).to_csv(OUT / "azure_ea_usage.csv", index=False)


def pc_invoice_recon() -> None:
    rows = []
    for cat, prod, sku, qty, price in [("Virtual Machines", "Azure plan", "D4s v5", 1440, 0.1632),
                                       ("Storage", "Azure plan", "Premium SSD", 30, 16.75),
                                       ("", "Microsoft 365 E3", "Microsoft 365 E3", 50, 36.0),
                                       ("", "Azure Reserved VM Instance", "D4s v5 1 Year", 2, 102.5)]:
        sub = round(qty * price, 2)
        rows.append({"PartnerId": "p1", "CustomerId": "c1", "CustomerName": "Fabrikam", "InvoiceNumber": "G012345678",
                     "ProductId": "DZH318Z0BQ4L" if cat else "CFQ7TTC0LF8S", "SkuName": sku, "ProductName": prod,
                     "MeterCategory": cat, "SubscriptionId": "s1", "ChargeStartDate": "9/1/2026",
                     "ChargeEndDate": "9/30/2026", "ChargeType": "new" if not cat else "usage",
                     "UnitPrice": price, "EffectiveUnitPrice": price, "Quantity": qty, "BillableQuantity": qty,
                     "Subtotal": sub, "TaxTotal": round(sub * 0.08, 2), "TotalForCustomer": round(sub * 1.08, 2),
                     "Currency": "USD"})
    pd.DataFrame(rows).to_csv(OUT / "pc_invoice_recon.csv", index=False)


def focus_parquet() -> None:
    rows = []
    for d in DAYS:
        for svc, cat, cost in [("Compute Engine", "Compute", 210.0), ("Cloud Storage", "Storage", 35.5),
                               ("BigQuery", "Analytics", 88.25), ("Cloud SQL", "Databases", 61.0)]:
            c = round(cost * rnd.uniform(0.9, 1.1), 4)
            rows.append({"BilledCost": c, "EffectiveCost": round(c * 0.97, 4), "ListCost": round(c * 1.1, 4),
                         "BillingCurrency": "USD", "ChargePeriodStart": f"{d}T00:00:00Z",
                         "ChargePeriodEnd": f"{d + timedelta(days=1)}T00:00:00Z", "ChargeCategory": "Usage",
                         "ProviderName": "Google Cloud", "ServiceName": svc, "ServiceCategory": cat,
                         "SubAccountId": "my-project", "RegionId": "us-central1", "ConsumedQuantity": 1.0,
                         "ConsumedUnit": "unit"})
    pd.DataFrame(rows).to_parquet(OUT / "focus_gcp.parquet", index=False)


def gcp_json() -> None:
    recs = []
    for d in DAYS[:10]:
        for sid, svc, sku, cost in [("6F81-5844-456A", "Compute Engine", "N2 Instance Core running in Americas", 120.0),
                                    ("95FF-2EF5-5EA1", "Cloud Storage", "Standard Storage US Multi-region", 12.0)]:
            recs.append({"billing_account_id": "01ABCD-234567-89EFGH",
                         "service": {"id": sid, "description": svc}, "sku": {"id": "AAAA-BBBB", "description": sku},
                         "usage_start_time": f"{d}T00:00:00Z", "usage_end_time": f"{d}T23:59:59Z",
                         "project": {"id": "shop-prod", "name": "shop-prod"}, "location": {"region": "us-east1"},
                         "cost": cost, "currency": "USD", "usage": {"amount_in_pricing_units": 24, "pricing_unit": "hour"},
                         "credits": [{"name": "Sustained use discount", "amount": -round(cost * 0.2, 2), "type": "SUSTAINED_USAGE_DISCOUNT"}],
                         "cost_type": "regular"})
    (OUT / "gcp_export.json").write_text(json.dumps(recs))


def cost_explorer() -> None:
    svcs = ["EC2-Instances($)", "Relational Database Service($)", "S3($)", "EC2-Other($)"]
    data = [[round(rnd.uniform(100, 400), 2) for _ in svcs] for _ in range(6)]
    months = ["2026-04-01", "2026-05-01", "2026-06-01", "2026-07-01", "2026-08-01", "2026-09-01"]
    lines = ["Service," + ",".join(svcs) + ",Total costs($)"]
    tot = [round(sum(r[i] for r in data), 2) for i in range(len(svcs))]
    lines.append("Service total," + ",".join(map(str, tot)) + f",{round(sum(tot), 2)}")
    for m, r in zip(months, data):
        lines.append(m + "," + ",".join(map(str, r)) + f",{round(sum(r), 2)}")
    (OUT / "aws_cost_explorer.csv").write_text("\n".join(lines))


def messy_csv() -> None:
    rows = [("Amazon EC2", "us-east-1", 4210.55), ("Amazon RDS", "us-east-1", 1820.10),
            ("Amazon S3", "us-east-1", 312.40), ("Amazon CloudFront", "global", 95.00)]
    total = round(sum(r[2] for r in rows), 2)
    text = "Monthly cloud spend report\nGenerated 2026-10-02 by FinOps team\n\n"
    text += "Service Name,Region,Monthly Cost (USD)\n"
    text += "\n".join(f"{a},{b},\"${c:,.2f}\"" for a, b, c in rows)
    text += f"\nTotal,,\"${total:,.2f}\"\n"
    (OUT / "messy_report.csv").write_text(text)


def custom_workbook() -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cloud Charges"
    ws["A1"] = "Northwind Managed Cloud - Monthly Statement"
    ws.merge_cells("A1:E1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws["B2"] = "Customer:", "Tailspin Toys"
    ws["A3"], ws["B3"] = "Period:", "September 2026"
    ws.append([])
    ws.append(["Item", "Details", "Qty", "Unit Price", "Amount"])
    for c in ws[5]:
        c.font = Font(bold=True)
    groups = {"Azure Compute": [("Virtual Machines", "D4s v5 - East US", 1440, 0.192),
                                ("Virtual Machines", "E8s v5 - East US", 720, 0.504),
                                ("Azure Kubernetes Service", "Uptime SLA", 720, 0.10)],
              "Azure Storage": [("Storage", "Premium SSD P30", 4, 135.17), ("Storage", "Blob Hot LRS 2 TB", 2048, 0.0184)],
              "Azure Databases": [("SQL Database", "General Purpose 4 vCore", 720, 0.505)]}
    r = 6
    for g, items in groups.items():
        ws.cell(r, 1, g).font = Font(bold=True)
        r += 1
        first = r
        for svc, det, q, p in items:
            ws.cell(r, 1, svc).alignment = Alignment(indent=1)
            ws.cell(r, 2, det)
            ws.cell(r, 3, q)
            ws.cell(r, 4, p)
            ws.cell(r, 5, f"=C{r}*D{r}")
            r += 1
        ws.cell(r, 4, "Subtotal").font = Font(bold=True)
        ws.cell(r, 5, f"=SUM(E{first}:E{r - 1})").font = Font(bold=True)
        r += 2
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Prices exclude tax. Contact billing@northwind.example for questions."
    _save_recalculated(wb, OUT / "custom_workbook.xlsx")


def pivot_workbook() -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spend by month"
    ws["A1"] = "Azure spend optimizer - 6 month view"
    ws["A1"].font = Font(bold=True)
    months = ["Apr 2026", "May 2026", "Jun 2026", "Jul 2026", "Aug 2026", "Sep 2026"]
    ws.append([])
    ws.append(["Service", *months, "Total"])
    svcs = {"Virtual Machines": 5200, "Storage": 900, "SQL Database": 1500, "Azure Kubernetes Service": 700, "Bandwidth": 210}
    r = 4
    for s, base in svcs.items():
        ws.cell(r, 1, s)
        for j in range(6):
            ws.cell(r, 2 + j, round(base * (1 + 0.02 * j) * rnd.uniform(0.97, 1.03), 2))
        ws.cell(r, 8, f"=SUM(B{r}:G{r})")
        r += 1
    ws.cell(r, 1, "Total").font = Font(bold=True)
    for j in range(7):
        col = "BCDEFGH"[j]
        ws.cell(r, 2 + j, f"=SUM({col}4:{col}{r - 1})")
    _save_recalculated(wb, OUT / "spend_optimizer_pivot.xlsx")


def _save_recalculated(wb, path: Path) -> None:
    """Real Excel files carry cached formula values; openpyxl doesn't write them. LibreOffice recalculates."""
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "src" / path.name
        src.parent.mkdir()
        wb.save(src)
        try:
            subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", d, str(src)],
                           check=True, capture_output=True, timeout=120)
            shutil.copy(Path(d) / path.name, path)
        except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
            wb.save(path)                                  # formulas without cached values


def reseller_pdf() -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(OUT / "reseller_invoice.pdf"), pagesize=letter)
    W, H = letter
    cols = [50, 330, 380, 440, 520]       # Description, Qty, Unit, Unit Price, Amount(right edge)

    def header(y):
        c.setFont("Helvetica-Bold", 9)
        for x, t in zip(cols, ["Description", "Qty", "Unit", "Unit Price", "Amount"]):
            c.drawString(x, y, t) if t != "Amount" else c.drawRightString(x + 40, y, t)
        return y - 16

    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, H - 50, "Contoso Cloud Partners LLC")
    c.setFont("Helvetica", 9)
    c.drawString(50, H - 66, "100 Main Street, Springfield")
    c.drawString(380, H - 50, "Invoice Number: INV-2026-0912")
    c.drawString(380, H - 64, "Billing Period: 09/01/2026 - 09/30/2026")
    c.drawString(380, H - 78, "Customer: Wide World Importers")
    y = header(H - 120)
    groups = {"Microsoft Azure": [("Virtual Machines D4s v5 East US", 1440, "Hours", 0.1920),
                                  ("Virtual Machines E8s v5 East US (reserved capacity not applied, billed at pay-as-you-go rates)", 720, "Hours", 0.5040),
                                  ("Premium SSD Managed Disks P30", 4, "Disks", 135.1700),
                                  ("Azure SQL Database General Purpose", 720, "Hours", 0.5050)],
              "Amazon Web Services": [("Amazon EC2 m5.xlarge", 2160, "Hours", 0.1920),
                                      ("Amazon RDS db.r6g.large", 720, "Hours", 0.2500),
                                      ("Amazon S3 Standard Storage", 5120, "GB-Mo", 0.0230)]}
    groups["Amazon Web Services"] += [(f"Amazon EC2 t3.medium instance {i}", 720, "Hours", 0.0416) for i in range(1, 31)]
    subtotal = 0.0
    for g, items in groups.items():
        c.setFont("Helvetica-Bold", 9)
        c.drawString(50, y, g)
        y -= 14
        for desc, q, unit, p in items:
            if y < 80:
                c.showPage()
                y = header(H - 60)
            amt = round(q * p, 2)
            subtotal += amt
            c.setFont("Helvetica", 9)
            first, rest = (desc[:55], desc[55:]) if len(desc) > 55 else (desc, "")
            c.drawString(60, y, first)
            c.drawRightString(cols[1] + 30, y, f"{q:,}")
            c.drawString(cols[2], y, unit)
            c.drawRightString(cols[3] + 50, y, f"${p:,.4f}")
            c.drawRightString(cols[4] + 40, y, f"${amt:,.2f}")
            y -= 12
            if rest:
                c.drawString(60, y, rest.strip())
                y -= 12
        y -= 6
    tax = round(subtotal * 0.0825, 2)
    c.setFont("Helvetica", 9)
    for label, val in [("Subtotal", subtotal), ("Sales Tax (8.25%)", tax)]:
        c.drawString(cols[3] - 60, y, label)
        c.drawRightString(cols[4] + 40, y, f"${val:,.2f}")
        y -= 14
    c.setFont("Helvetica-Bold", 10)
    c.drawString(cols[3] - 60, y, "Total Due")
    c.drawRightString(cols[4] + 40, y, f"${subtotal + tax:,.2f}")
    c.save()


def list_pdf() -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(OUT / "summary_invoice.pdf"), pagesize=letter)
    W, H = letter
    c.setFont("Helvetica-Bold", 13)
    c.drawString(50, H - 50, "Blue Yonder IT Services - Cloud Usage Summary")
    c.setFont("Helvetica", 10)
    c.drawString(50, H - 70, "Statement date: October 3, 2026")
    items = [("Amazon EC2", 8423.17), ("Amazon RDS", 2210.40), ("Amazon S3", 512.88), ("AWS Lambda", 41.07),
             ("Amazon CloudFront", 188.10), ("AWS Support (Business)", 1137.56)]
    y = H - 110
    for name, amt in items:
        dots = "." * max(3, 60 - len(name))
        c.drawString(50, y, f"{name} {dots} ${amt:,.2f}")
        y -= 16
    y -= 10
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, f"Total Due: ${sum(a for _, a in items):,.2f}")
    c.save()


def custom_json() -> None:
    lines = [{"svc": "Virtual Machines", "detail": "D4s v5", "qty": "1440", "amt": "$276.48"},
             {"svc": "Storage", "detail": "Blob Hot", "qty": "2048", "amt": "$37.68"},
             {"svc": "SQL Database", "detail": "GP 4 vCore", "qty": "720", "amt": "$363.60"}]
    doc = {"customer": {"name": "Litware"}, "invoice": {"number": "LW-77", "lines": lines, "grand": "$677.76"}}
    (OUT / "custom_api.json").write_text(json.dumps(doc, indent=1))


def german_xlsx(path: Path, month: str = "September 2026", amount_header: str = "Betrag", scale: float = 1.0) -> None:
    """EU reseller invoice: German headers, decimal commas, total row. The rule mapper can't read it."""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rechnung"
    ws["A1"] = "Fabrikam Cloud GmbH - Rechnung"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = f"Leistungszeitraum: {month}"
    ws.append([])
    ws.append(["Leistung", "Menge", "Einzelpreis", amount_header])
    items = [("Virtual Machines D4s v5", 1440, 0.1920), ("Azure SQL Database", 720, 0.5050),
             ("Storage Premium SSD", 4, 135.1700), ("Bandwidth", 512, 0.0870)]
    total = 0.0
    for name, q, p in items:
        amt = round(q * p * scale, 2)
        total += amt
        ws.append([name, q, f"{p * scale:.4f}".replace(".", ","), f"{amt:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")])
    ws.append(["Gesamtbetrag", None, None, f"{total:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    wb.save(path)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for fn in (aws_cur2, azure_ea, pc_invoice_recon, focus_parquet, gcp_json, cost_explorer, messy_csv,
               custom_workbook, pivot_workbook, reseller_pdf, list_pdf, custom_json):
        fn()
        print("ok", fn.__name__)
    for fn in (aws_cur2_commitments, aws_cur2_no_amortized, azure_amortized, azure_actual, focus_commitments,
               gcp_commitments, aws_cur_legacy, aws_cur2_product_map, aws_cur_legacy_full, aws_cur2_flat,
               aws_cur2_map, azure_usage_full, focus_full, gcp_export_full):
        print("ok", fn.__name__, "total", fn())
    german_xlsx(OUT / "german_invoice.xlsx")
    print("ok german_xlsx")
