"""The document model's readers: what a file becomes before any spec reads it."""

from __future__ import annotations

import io

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from billparse import docmodel


def _workbook(rows: int, cols: int = 10) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([f"Col {c}" for c in range(cols)])
    for r in range(rows):
        ws.append([f"r{r}c{c}" if c % 3 else r * 1.5 for c in range(cols)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_a_sheet_asks_its_size_once_not_per_row(monkeypatch):
    """Worksheet.max_column scans every cell: asked per row, a sheet of 11,000 rows took minutes (rows x cells)."""
    calls = {"n": 0}
    real = Worksheet.max_column

    def counted(self):
        calls["n"] += 1
        return real.fget(self)
    monkeypatch.setattr(Worksheet, "max_column", property(counted))
    doc = docmodel.load_bytes(_workbook(3000), "big.xlsx")
    sheet = next(r for r in doc.regions if r.kind == "sheet")
    assert len(sheet.rows) == 3001 and len(sheet.rows[0].cells) == 10
    assert calls["n"] <= 5


def _usage_book(months: dict[str, int], total_row: bool = False, reservations: bool = False) -> bytes:
    """A reseller usage export: one sheet per month, the same 8 columns; some meters and resources have "Total" in
    their name (real Azure names do)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head = ["Customer", "Subscription Name", "Resource Name", "Meter Name", "Consumed Qty", "Unit", "MSRP",
            "Extended MSRP"]
    for title, n in months.items():
        ws = wb.create_sheet(title)
        ws.append(head)
        for i in range(n):
            meter = "Total Storage Transactions" if i % 7 == 3 else "D4s v5"
            resource = "vm-total-backup" if i % 11 == 5 else f"vm-{i}"
            ws.append(["Contoso", "Prod", resource, meter, 24.0, "1 Hour", 0.192, round(24 * 0.192, 2)])
        if total_row:
            ws.append([None, None, None, None, None, None, "Total", round(n * 24 * 0.192, 2)])
    if reservations:                                   # another table: not the usage table's continuation
        ws = wb.create_sheet("Reservations")
        ws.append(["Name", "Reservation Id", "Status", "Term", "Region", "Product name", "Quantity", "Utilization %"])
        for i in range(5):
            ws.append([f"res-{i}", f"id-{i}", "Succeeded", "P1Y", "eastus", "Standard_D4s_v5", 2, 97.5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_a_data_row_named_total_does_not_end_the_table():
    """A line whose meter or resource name says "Total" is a line, not the file's total: the reading went on to the
    end of every sheet, and a real total row (a label and an amount) still ends its table."""
    from billparse import mapper
    from billparse import spec as spec_mod
    doc = docmodel.load_bytes(_usage_book({"May": 30, "June": 25}, reservations=True), "usage.xlsx")
    spec, _, _ = mapper.propose_mapped(doc)
    ext = spec_mod.run(spec, doc)
    assert len(ext.rows) == 55 and not ext.rows["_src"].str.startswith("Reservations").any()
    assert not [t for t in ext.stated_totals if t.get("kind") == "total"]
    doc = docmodel.load_bytes(_usage_book({"May": 30}, total_row=True), "usage.xlsx")
    ext = spec_mod.run(spec, doc)
    assert len(ext.rows) == 30 and [t["kind"] for t in ext.stated_totals] == ["total"]
